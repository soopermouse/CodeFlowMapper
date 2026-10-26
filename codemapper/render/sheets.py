"""Excel workbooks: Flow tracker, Database tables and calls, Tooling and dependencies."""
from __future__ import annotations
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from ..textutil import short, uniq
from .covers import nature

HEAD_FILL = PatternFill('solid', fgColor='FF1F3864')
HEAD_FONT = Font(name='Arial', size=10, bold=True, color='FFFFFFFF')
BODY = Font(name='Arial', size=10)
THIN = Side(style='thin', color='FFBFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ZEBRA = PatternFill('solid', fgColor='FFF4F6F8')


def sheet(wb, name, header, rows, widths=None):
    ws = wb.create_sheet(name[:31])
    ws.append(header)
    for r in rows:
        ws.append([_cell(v) for v in r])
    for j in range(1, len(header) + 1):
        c = ws.cell(row=1, column=j)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(j)].width = (widths[j - 1] if widths and j - 1 < len(widths) else 18)
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for c in row:
            c.font, c.border = BODY, BORDER
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if i % 2 == 1:
                c.fill = ZEBRA
    ws.freeze_panes = 'B2'
    if rows:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(header))}{len(rows) + 1}'
    return ws


def _cell(v):
    if isinstance(v, (list, tuple, set)):
        v = '\n'.join(str(x) for x in v)
    if isinstance(v, str) and len(v) > 32000:
        v = v[:32000]
    return v


def _book(author, title):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.properties.creator = author or ''
    wb.properties.lastModifiedBy = author or ''
    wb.properties.title = title
    return wb


def data_sources(f):
    out = []
    rw = defaultdict(set)
    for s in f.sql:
        for t in s.tables:
            rw[t].add('write' if s.kind not in ('SELECT', 'WITH', 'SQL') else 'read')
    if rw:
        out.append('Tables: ' + ', '.join(f'{t} ({"/".join(sorted(v))})' for t, v in rw.items()))
    if f.procedures:
        out.append('Stored procedures: ' + ', '.join(f.procedures))
    if f.middle_layer:
        out.append('Middle layer: ' + ', '.join(uniq(c.name + (f" '{c.target}'" if c.target else '') for c in f.middle_layer)))
    if f.connections:
        out.append('Connections: ' + ', '.join(f.connections))
    return '\n'.join(out) or 'None (no data access)'


def flow_tracker(project, path, cfg):
    ids = {f.id: f for f in project.flows}
    tools = {t.id: t for t in project.tools}
    wb = _book(cfg.get('author'), 'Flow tracker')
    rows = []
    for f in project.flows:
        e = f.entry
        upf = sorted({ids[u].functionality for u in f.upstream if ids[u].functionality != f.functionality})
        dnf = sorted({ids[d].functionality for d in f.downstream if ids[d].functionality != f.functionality})
        rows.append([f.id, f.name, f.business, f.functionality, f.module, f.submodule, e.kind.replace('_', ' '), e.route if e.kind != 'page' else e.file,
                     e.method, [f'{u} {ids[u].name}' for u in f.upstream] or '-', upf or ('Same functionality' if f.upstream else '-'),
                     [f'{d} {ids[d].name}' for d in f.downstream] or '-', dnf or ('Same functionality' if f.downstream else '-'),
                     data_sources(f), nature(f), [f'{t} {tools[t].name}' for t in f.tools], [f'{t} {tools[t].name}' for t in f.dependencies],
                     ', '.join(f.roles) or '-', {'code': 'Code', 'ocr': 'Photos', 'sonar': 'SonarQube'}.get(f.origin, f.origin)])
    sheet(wb, 'Flow tracker', ['Flow number', 'Flow name', 'Business name', 'Functionality', 'Module', 'Submodule', 'Entry kind', 'Entry point', 'Method',
                               'Upstream flows', 'Upstream functionalities', 'Downstream flows', 'Downstream functionalities', 'Data sources',
                               'Nature', 'Tools', 'Dependencies', 'Access (roles / rights)', 'Source'], rows,
          [11, 34, 28, 30, 14, 18, 14, 40, 9, 34, 30, 34, 30, 60, 16, 34, 34, 24, 10])
    det = []
    for f in project.flows:
        e = f.entry
        det.append([f.id, f.name, e.file, e.handler.split('@')[0] if e.handler else '-', ', '.join(f'{n} ({s})' for s, n in f.params) or '-',
                    len(f.forms), sum(len(fm.fields) for fm in f.forms), sum(1 for fm in f.forms for x in fm.fields if x.kind == 'select'),
                    len(f.branch_detail), ', '.join(f.templates) or '-', '\n'.join(f.files[1:]) or '-', '\n'.join(k.split('@')[0] for k in f.functions) or '-',
                    len(f.calls), len(f.sql), ', '.join(f.tables) or '-', ', '.join(f.procedures) or '-', len(f.messages),
                    ', '.join(f'{k} {n}' if k == 'template' else k for k, n, _ in f.outputs) or '-', len(f.issues)])
    sheet(wb, 'Flow details', ['Flow number', 'Flow name', 'Entry file', 'Handler', 'Request input', 'Forms', 'Fields', 'Dropdowns', 'Decisions',
                               'Templates', 'Files', 'Functions reached', 'Calls', 'SQL statements', 'Tables', 'Stored procedures', 'Messages',
                               'Outputs', 'SonarQube issues'], det, [11, 30, 40, 30, 30, 7, 7, 9, 9, 30, 40, 40, 7, 9, 30, 30, 9, 24, 9])
    fr = []
    for f in project.flows:
        for fm in f.forms:
            for x in fm.fields:
                fr.append([f.id, fm.name, fm.method, fm.action, x.name, x.kind, x.label, 'Yes' if x.required else '', x.maxlength, x.pattern,
                           '\n'.join(f'{o.value} = {o.label}' + (' (default)' if o.selected else '') for o in x.options), x.options_source, fm.file])
    sheet(wb, 'Fields and dropdowns', ['Flow number', 'Form', 'Method', 'Action', 'Field', 'Type', 'Label', 'Required', 'Max length', 'Pattern',
                                       'Options (value = label)', 'Options from code', 'File'], fr, [11, 20, 8, 24, 20, 10, 24, 9, 9, 16, 40, 36, 36])
    dr = []
    for f in project.flows:
        for b in f.branch_detail:
            for cs in b['cases']:
                dr.append([f.id, f'{b["kind"]} ({b["subject"]})', cs['label'], '\n'.join(f'{n}({a})' for n, a, k in cs['calls']),
                           '\n'.join(f'{k} {p or ", ".join(t)}' for k, t, p in cs['sql']), '\n'.join(cs['messages']), '\n'.join(cs['links'])])
    sheet(wb, 'Decisions', ['Flow number', 'Decision', 'Case', 'Calls', 'SQL', 'Messages', 'Goes to'], dr, [11, 34, 20, 50, 34, 34, 30])
    cr = []
    for f in project.flows:
        for d, c, k in f.calls:
            fn = project.functions.get(k)
            cr.append([f.id, d, c.name, c.args, k.split('@')[0] if k else 'library / framework', f'{fn.file}:{fn.line}' if fn else f'{c.file}:{c.line}'])
    sheet(wb, 'Calls', ['Flow number', 'Level', 'Call', 'Arguments', 'Resolved to', 'Where'], cr, [11, 6, 36, 50, 36, 40])
    ed = [[a, ids[a].name, b, ids[b].name, kind, via] for a, b, kind, via in project.edges]
    sheet(wb, 'Flow dependencies', ['Upstream flow', 'Name', 'Downstream flow', 'Name ', 'How', 'Through'], ed, [12, 34, 14, 34, 14, 40])
    fu = defaultdict(list)
    for f in project.flows:
        fu[f.functionality].append(f)
    furows = []
    for n, fl in fu.items():
        mine = {f.id for f in fl}
        furows.append([n, fl[0].module, fl[0].submodule, ', '.join(f.id for f in fl), len(fl),
                       sorted({ids[u].functionality for f in fl for u in f.upstream if u not in mine}) or '-',
                       sorted({ids[d].functionality for f in fl for d in f.downstream if d not in mine}) or '-'])
    sheet(wb, 'Functionalities', ['Functionality', 'Module', 'Submodule', 'Flows', 'Number of flows', 'Upstream functionalities', 'Downstream functionalities'],
          furows, [34, 16, 20, 40, 10, 36, 36])
    if any(f.issues for f in project.flows):
        ir = []
        for f in project.flows:
            for i in f.issues:
                ir.append([f.id, i.get('file'), i.get('line'), i.get('severity'), i.get('type'), i.get('rule'), i.get('message')])
        sheet(wb, 'SonarQube issues', ['Flow number', 'File', 'Line', 'Severity', 'Type', 'Rule', 'Message'], ir, [11, 40, 7, 11, 14, 24, 60])
    wb.save(path)


def database(project, path, cfg):
    wb = _book(cfg.get('author'), 'Database tables and calls')
    tb = defaultdict(lambda: defaultdict(set))
    for f in project.flows:
        for s in f.sql:
            for t in s.tables:
                tb[t][f.id].add('write' if s.kind not in ('SELECT', 'WITH', 'SQL') else 'read')
    rows = [[t, len(fl), '\n'.join(f'{k} ({"/".join(sorted(v))})' for k, v in sorted(fl.items()))] for t, fl in sorted(tb.items(), key=lambda kv: kv[0].lower())]
    sheet(wb, 'Tables', ['Table', 'Number of flows', 'Flows (access)'], rows, [36, 12, 60])
    pr = defaultdict(set)
    for f in project.flows:
        for p in f.procedures:
            pr[p].add(f.id)
    sheet(wb, 'Stored procedures', ['Stored procedure', 'Number of flows', 'Flows'], [[p, len(v), ', '.join(sorted(v))] for p, v in sorted(pr.items())], [40, 12, 60])
    sr = []
    for f in project.flows:
        for s in f.sql:
            sr.append([f.id, f.name, s.kind, ', '.join(s.tables), s.procedure, ', '.join(s.params), s.text, s.function or '-', f'{s.file}:{s.line}'])
    sheet(wb, 'Calls per flow', ['Flow number', 'Flow name', 'Statement', 'Tables', 'Stored procedure', 'Parameters', 'SQL', 'Function', 'Where'], sr,
          [11, 30, 10, 30, 24, 20, 70, 30, 36])
    mr = []
    for f in project.flows:
        for c in f.middle_layer:
            mr.append([f.id, f.name, c.name, c.target, c.args, f'{c.file}:{c.line}'])
    sheet(wb, 'Middle-layer calls', ['Flow number', 'Flow name', 'Mechanism', 'Target', 'Arguments', 'Where'], mr, [11, 30, 30, 34, 50, 36])
    cn = defaultdict(set)
    for f in project.flows:
        for c in f.connections:
            cn[c].add(f.id)
    sheet(wb, 'Connections', ['Connection', 'Flows'], [[c, ', '.join(sorted(v))] for c, v in sorted(cn.items())], [40, 60])
    wb.save(path)


def tooling(project, path, cfg):
    wb = _book(cfg.get('author'), 'Tooling and dependencies')
    mn = int(cfg.get('tool_min_flows', 3))
    tr = [[t.id, t.name, t.kind, '\n'.join(t.functions[:60]), len(t.flows), ', '.join(t.flows), '\n'.join(t.functionalities)]
          for t in project.tools if t.kind != 'Dependency']
    sheet(wb, 'Tools', ['Tool', 'Function', 'Type', 'Legacy functions', 'Number of flows', 'Flows', 'Functionalities'], tr, [8, 34, 13, 50, 10, 40, 40])
    dr = [[t.id, t.name, 'function group', '\n'.join(t.functions[:60]), len(t.flows), ', '.join(t.flows), '\n'.join(t.functionalities)]
          for t in project.tools if t.kind == 'Dependency']
    dr += [['-', d['file'], d['kind'], '-', len(d['flows']), ', '.join(d['flows']), '\n'.join(d['functionalities'])]
           for d in project.file_dependencies if d['type'] == 'Dependency']
    sheet(wb, 'Dependencies', ['Id', 'Dependency', 'Kind', 'Legacy functions', 'Number of flows', 'Flows', 'Functionalities'], dr, [8, 40, 14, 46, 10, 30, 40])
    sh = [[d['file'], d['kind'], len(d['flows']), ', '.join(d['flows']), '\n'.join(d['functionalities'])] for d in sorted(project.file_dependencies, key=lambda d: -len(d['flows'])) if d['type'] != 'Dependency']
    sheet(wb, 'Shared files', ['File', 'Kind', 'Number of flows', 'Flows', 'Functionalities'], sh, [44, 14, 10, 50, 44])
    mods = defaultdict(list)
    for f in project.flows:
        mods[f.module].append(f)
    tools = {t.id: t for t in project.tools}
    mr = []
    for m, fl in sorted(mods.items()):
        ts = sorted({t for f in fl for t in f.tools})
        ds = sorted({t for f in fl for t in f.dependencies})
        mr.append([m, len({f.functionality for f in fl}), len(fl), '\n'.join(f'{t} {tools[t].name}' for t in ts) or '-', '\n'.join(f'{t} {tools[t].name}' for t in ds) or 'None'])
    sheet(wb, 'Tools by module', ['Module', 'Functionalities', 'Flows', 'Tools', f'Dependencies (fewer than {mn} flows)'], mr, [20, 12, 8, 50, 50])
    fu = defaultdict(list)
    for f in project.flows:
        fu[f.functionality].append(f)
    fr = [[n, ', '.join(f.id for f in fl), '\n'.join(f'{t} {tools[t].name}' for t in sorted({t for f in fl for t in f.tools})) or '-',
           '\n'.join(f'{t} {tools[t].name}' for t in sorted({t for f in fl for t in f.dependencies})) or '-'] for n, fl in sorted(fu.items())]
    sheet(wb, 'Tools by functionality', ['Functionality', 'Flows', 'Tools', 'Dependencies'], fr, [34, 30, 50, 50])
    fd = [[f.id, f.name, p, next((d['kind'] for d in project.file_dependencies if d['file'] == p), ''), next((d['type'] for d in project.file_dependencies if d['file'] == p), '')]
          for f in project.flows for p in f.files[1:]]
    sheet(wb, 'File dependencies', ['Flow number', 'Flow name', 'File', 'Kind', 'Type'], fd, [11, 34, 50, 14, 12])
    wb.save(path)
